"""Block model and the Markdown / DOCX renderers for the final project report.

The report text lives once, in final_report_content.py, as a list of blocks.  This module numbers the
headings, figures, tables, equations and citations, then writes the same blocks to Markdown and to an
editable Word document.  The PDF is exported from that Word document (see generate_final_report.py).
"""
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

ACCENT_HEX, INK_HEX, MUTED_HEX = "4338A8", "25232B", "6B6760"
TINT_HEX, LINE_HEX = "ECEAFB", "C9C4DD"
BODY_FONT, HEAD_FONT, MONO_FONT = "Times New Roman", "Arial", "Consolas"
FRONT_HEADINGS = ("Table of Contents", "List of Figures", "List of Tables")


# ====================================================================== block constructors
def H1(title, num=True, page_break=False):
    return {"t": "h1", "title": title, "num": num, "pb": page_break}


def H2(title):
    return {"t": "h2", "title": title}


def H3(title):
    return {"t": "h3", "title": title}


def P(text):
    return {"t": "p", "text": text}


def BUL(*items):
    return {"t": "bul", "items": list(items)}


def NUM(*items):
    return {"t": "num", "items": list(items)}


def EQ(key, tex):
    return {"t": "eq", "key": key, "tex": tex}


def CASES(key, lhs, rows):
    """Piecewise definition: rows = [(expression_tex, condition_tex), ...]."""
    return {"t": "cases", "key": key, "lhs": lhs, "rows": rows}


def FIG(key, path, caption, width=6.0):
    return {"t": "fig", "key": key, "path": path, "caption": caption, "width": width}


def TAB(key, caption, header, rows, widths, align=None, size=9.0, mono_cols=()):
    return {"t": "tab", "key": key, "caption": caption, "header": header, "rows": rows,
            "widths": widths, "align": align or ["l"] * len(header), "size": size, "mono": mono_cols}


def CODE(text):
    return {"t": "code", "text": text}


def NOTE(title, text):
    return {"t": "note", "title": title, "text": text}


def REFS(items):
    return {"t": "refs", "items": items}


# ====================================================================== numbering pass
class Numbering:
    """Assigns heading numbers, figure/table/equation numbers and citation numbers."""

    def __init__(self, blocks, refs):
        self.refs = refs                      # key -> formatted reference text
        self.cite_order = []
        self.used = set()                    # (kind, key) of every figure / table / equation reference
        self.labels = {}                      # ("fig", key) -> "6.1"
        self.headings = []                    # (level, label, title, block)
        chapter = None
        sub = 0
        counters = {}
        letters = iter("ABCDEFGHIJ")
        n = 0
        for b in blocks:
            t = b["t"]
            if t == "h1":
                if b["num"] is True:
                    n += 1
                    chapter, b["label"] = str(n), f"{n}."
                elif b["num"] == "appendix":
                    chapter = next(letters)
                    b["label"] = f"Appendix {chapter}:"
                else:
                    chapter, b["label"] = None, ""
                sub = 0
                self.headings.append((1, b["label"], b["title"], b))
            elif t == "h2":
                sub += 1
                b["label"] = f"{chapter}.{sub}" if chapter else ""
                self.headings.append((2, b["label"], b["title"], b))
            elif t in ("fig", "tab", "eq", "cases"):
                kind = "fig" if t == "fig" else "tab" if t == "tab" else "eq"
                counters[(kind, chapter)] = counters.get((kind, chapter), 0) + 1
                b["label"] = f"{chapter}.{counters[(kind, chapter)]}"
                self.labels[(kind, b["key"])] = b["label"]
        self._blocks = blocks

    def resolve(self, text):
        def sub(m):
            kind, key = m.group(1), m.group(2)
            if kind == "cite":
                nums = []
                for k in key.split(","):
                    if k not in self.refs:
                        raise KeyError(f"unknown citation {k}")
                    if k not in self.cite_order:
                        self.cite_order.append(k)
                    nums.append(f"[{self.cite_order.index(k) + 1}]")
                return ", ".join(nums)
            label = self.labels[(kind, key)]
            self.used.add((kind, key))
            return {"fig": f"Figure {label}", "tab": f"Table {label}", "eq": f"Eq. ({label})"}[kind]
        return re.sub(r"\[\[(\w+):([\w,]+)\]\]", sub, text)

    def ordered_refs(self):
        return [(i + 1, k, self.refs[k]) for i, k in enumerate(self.cite_order)]


def prepare(blocks, refs):
    """Number everything and resolve cross references in place.  Citation order = order of first use."""
    num = Numbering(blocks, refs)

    def fix(v):
        return num.resolve(v) if isinstance(v, str) else v

    for b in blocks:
        for k in ("text", "title", "caption"):
            if k in b and isinstance(b[k], str):
                b[k] = num.resolve(b[k])
        if "items" in b and b["t"] in ("bul", "num"):
            b["items"] = [num.resolve(i) for i in b["items"]]
        if b["t"] == "tab":
            b["header"] = [fix(h) for h in b["header"]]
            b["rows"] = [[fix(c) for c in row] for row in b["rows"]]
    uncited = [k for k in num.labels if k[0] in ("fig", "tab") and k not in num.used]
    if uncited:
        raise ValueError(f"figures / tables never referenced in the text: {uncited}")
    unused = set(refs) - set(num.cite_order)
    if unused:
        raise ValueError(f"references never cited: {sorted(unused)}")
    for b in blocks:
        if b["t"] == "refs":
            b["items"] = num.ordered_refs()
    return num


# ====================================================================== equations as images
def render_equation(tex, path, size=13):
    plt.rcParams["mathtext.fontset"] = "stix"
    fig = plt.figure(figsize=(0.1, 0.1))
    fig.text(0, 0, f"${tex}$", fontsize=size, color="#" + INK_HEX)
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.04, transparent=True)
    plt.close(fig)


def _brace(ax, x, y0, y1, w=0.1):
    """Left curly brace spanning y0..y1 with its tip at x - w."""
    from matplotlib.patches import PathPatch
    from matplotlib.path import Path
    ym = (y0 + y1) / 2
    M, L, C = Path.MOVETO, Path.LINETO, Path.CURVE3
    verts = [(x + w, y1), (x, y1), (x, y1 - w), (x, ym + w), (x, ym), (x - w, ym), (x, ym), (x, ym - w),
             (x, y0 + w), (x, y0), (x + w, y0)]
    codes = [M, C, C, L, C, C, C, C, L, C, C]
    ax.add_patch(PathPatch(Path(verts, codes), fc="none", ec="#" + INK_HEX, lw=1.1))


def render_cases(lhs, rows, path, size=13):
    """Piecewise function: left-hand side, a curly brace and aligned expression / condition columns."""
    plt.rcParams["mathtext.fontset"] = "stix"
    probe = plt.figure(figsize=(8, 2))
    renderer = probe.canvas.get_renderer()

    def width(s):
        t = probe.text(0, 0, f"${s}$", fontsize=size)
        w = t.get_window_extent(renderer).width / probe.dpi
        t.remove()
        return w

    heights = [0.52 if "frac" in expr else 0.34 for expr, _ in rows]
    total = sum(heights) + 0.06
    wl = width(lhs) + width("=") + 0.1
    x_brace = wl + 0.18
    x_expr = x_brace + 0.22
    x_cond = x_expr + max(width(e) for e, _ in rows) + 0.5
    total_w = x_cond + max(width(c) for _, c in rows) + 0.05
    plt.close(probe)
    fig = plt.figure(figsize=(total_w, total))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.axis("off")
    ax.set_xlim(0, total_w)
    ax.set_ylim(0, total)
    ax.text(0.0, total / 2, f"${lhs}=$", fontsize=size, va="center", color="#" + INK_HEX)
    _brace(ax, x_brace, 0.04, total - 0.04)
    y = total - 0.03
    for (expr, cond), h in zip(rows, heights):
        ax.text(x_expr, y - h / 2, f"${expr}$", fontsize=size, va="center", color="#" + INK_HEX)
        ax.text(x_cond, y - h / 2, f"${cond}$", fontsize=size, va="center", color="#" + INK_HEX)
        y -= h
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.04, transparent=True)
    plt.close(fig)


# ====================================================================== inline markup
TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|_\{[^}]*\}|\^\{[^}]*\})")


def parse_inline(text):
    """Split inline markup into (kind, text) runs: plain, b, i, code, sub, sup."""
    out = []
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            out.append(("plain", text[pos:m.start()]))
        tok = m.group(0)
        if tok.startswith("**"):
            out.append(("b", tok[2:-2]))
        elif tok.startswith("`"):
            out.append(("code", tok[1:-1]))
        elif tok.startswith("_{"):
            out.append(("sub", tok[2:-1]))
        elif tok.startswith("^{"):
            out.append(("sup", tok[2:-1]))
        else:
            out.append(("i", tok[1:-1]))
        pos = m.end()
    if pos < len(text):
        out.append(("plain", text[pos:]))
    return out


# ====================================================================== Markdown
def _md_inline(text):
    parts = []
    for kind, s in parse_inline(text):
        parts.append({"plain": s, "b": f"**{s}**", "i": f"*{s}*", "code": f"`{s}`",
                      "sub": f"<sub>{s}</sub>", "sup": f"<sup>{s}</sup>"}[kind])
    return "".join(parts)


def _md_cell(text):
    return _md_inline(text).replace("|", "\\|").replace("\n", "<br>")


def write_markdown(blocks, ctx, path, assets_rel="report_assets"):
    cover = ctx["cover"]
    out = [f"# {cover['title']}", "",
           f"**Course:** {cover['course']}", "", "**Submitted by:**", ""]
    out += [f"- {n} — Registration Number {r}" for n, r in cover["students"]]
    out += ["", f"**Submitted to:** {cover['faculty']}", f"**Department / programme / institution:** {cover['institution']}",
            f"**Date of submission:** {cover['date']}", "", "---", ""]
    heads = [(lvl, lab, title) for lvl, lab, title, _ in ctx["numbering"].headings]
    out += ["## Table of Contents", ""]
    for lvl, lab, title in heads:
        if lvl == 1:
            out.append(f"- **{(lab + ' ') if lab else ''}{title}**")
        else:
            out.append(f"  - {lab} {title}")
    out += ["", "## List of Figures", ""]
    out += [f"- Figure {b['label']}: {_md_inline(b['caption'])}" for b in blocks if b["t"] == "fig"]
    out += ["", "## List of Tables", ""]
    out += [f"- Table {b['label']}: {_md_inline(b['caption'])}" for b in blocks if b["t"] == "tab"]
    out += ["", "---", ""]
    for b in blocks:
        t = b["t"]
        if t == "h1":
            out += [f"## {(b['label'] + ' ') if b['label'] else ''}{b['title']}", ""]
        elif t == "h2":
            out += [f"### {b['label']} {b['title']}", ""]
        elif t == "h3":
            out += [f"#### {b['title']}", ""]
        elif t == "p":
            out += [_md_inline(b["text"]), ""]
        elif t == "bul":
            out += [f"- {_md_inline(i)}" for i in b["items"]] + [""]
        elif t == "num":
            out += [f"{k}. {_md_inline(i)}" for k, i in enumerate(b["items"], 1)] + [""]
        elif t == "eq":
            out += ["$$", f"{b['tex']} \\qquad \\text{{({b['label']})}}", "$$", ""]
        elif t == "cases":
            body = " \\\\\n".join(f"{e} & {c}" for e, c in b["rows"])
            out += ["$$", f"{b['lhs']} = \\begin{{cases}}", body, f"\\end{{cases}} \\qquad \\text{{({b['label']})}}", "$$", ""]
        elif t == "fig":
            out += [f"![Figure {b['label']}]({assets_rel}/{os.path.basename(b['path'])})", "",
                    f"*Figure {b['label']}: {_md_inline(b['caption'])}*", ""]
        elif t == "tab":
            out += [f"**Table {b['label']}: {_md_inline(b['caption'])}**", "",
                    "| " + " | ".join(_md_cell(h) for h in b["header"]) + " |",
                    "|" + "|".join(" --- " for _ in b["header"]) + "|"]
            out += ["| " + " | ".join(_md_cell(c) for c in row) + " |" for row in b["rows"]] + [""]
        elif t == "code":
            out += ["```text", b["text"].rstrip("\n"), "```", ""]
        elif t == "note":
            out += [f"> **{b['title']}** {_md_inline(b['text'])}", ""]
        elif t == "refs":
            out += [f"{n}. {_md_inline(text)}" for n, _, text in b["items"]] + [""]
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(out).rstrip() + "\n")


# ====================================================================== DOCX helpers

PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight",
              "tblHeader", "tblCellSpacing", "jc", "hidden"]
SECTPR_ORDER = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
                "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
                "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings"]


def _put(parent, child, order):
    """Insert child into parent at the position the OOXML schema requires (Word rejects wrong order)."""
    name = child.tag.split("}")[1]
    for old in parent.findall(child.tag):
        parent.remove(old)
    idx = order.index(name)
    for i, existing in enumerate(list(parent)):
        ename = existing.tag.split("}")[1]
        if ename in order and order.index(ename) > idx:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child

def _rfonts(el, name):
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)


def _style(styles, name, font, size, bold=None, italic=None, color=None, base=None):
    try:
        st = styles[name]
    except KeyError:
        from docx.enum.style import WD_STYLE_TYPE
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles[base or "Normal"]
    st.font.name = font
    st.font.size = Pt(size)
    if bold is not None:
        st.font.bold = bold
    if italic is not None:
        st.font.italic = italic
    if color:
        st.font.color.rgb = RGBColor.from_string(color)
    _rfonts(st.element, font)
    return st


def _shd(fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    return shd


def _cell_shade(cell, fill):
    _put(cell._tc.get_or_add_tcPr(), _shd(fill), TCPR_ORDER)


def _para_shade(p, fill):
    _put(p._p.get_or_add_pPr(), _shd(fill), PPR_ORDER)


def _cell_borders(cell, **edges):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = edges.get(edge)
        el = OxmlElement(f"w:{edge}")
        if spec:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), spec[1])
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    _put(tcpr, borders, TCPR_ORDER)


def _cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    _put(tblpr, mar, TBLPR_ORDER)


def _field(paragraph, instr, placeholder="", dirty=False):
    """Insert a complex field (begin / instruction / separate / cached result / end)."""
    def fc(kind):
        r = paragraph.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        if dirty and kind == "begin":
            el.set(qn("w:dirty"), "true")
        r._r.append(el)
        return r
    fc("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fc("separate")
    paragraph.add_run(placeholder)
    fc("end")


def _add_runs(paragraph, text, base_size=None, font=None, color=None):
    for kind, s in parse_inline(text):
        run = paragraph.add_run(s)
        if kind == "b":
            run.bold = True
        elif kind == "i":
            run.italic = True
        elif kind == "code":
            run.font.name = MONO_FONT
            _rfonts(run._r, MONO_FONT)
            run.font.size = Pt((base_size or 11) - 1.5)
        elif kind == "sub":
            run.font.subscript = True
        elif kind == "sup":
            run.font.superscript = True
        if font and kind != "code":
            run.font.name = font
            _rfonts(run._r, font)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if base_size and kind != "code":
            run.font.size = Pt(base_size)


def _bottom_border(paragraph, color=ACCENT_HEX, sz=8, space=1):
    ppr = paragraph._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), color)
    pb.append(b)
    _put(ppr, pb, PPR_ORDER)


def _keep_next(p, on=True):
    p.paragraph_format.keep_with_next = on


# ====================================================================== DOCX writer
def _setup_styles(doc):
    st = doc.styles
    n = _style(st, "Normal", BODY_FONT, 10.5, color=INK_HEX)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.12
    n.paragraph_format.widow_control = True
    h1 = _style(st, "Heading 1", HEAD_FONT, 17, bold=True, color=ACCENT_HEX)
    h1.paragraph_format.space_before = Pt(20)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True
    h2 = _style(st, "Heading 2", HEAD_FONT, 13, bold=True, color=INK_HEX)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True
    h3 = _style(st, "Heading 3", HEAD_FONT, 11, bold=True, italic=True, color=INK_HEX)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True
    for name, size, italic in (("FigCaption", 9.5, True), ("TabCaption", 9.5, True)):
        c = _style(st, name, HEAD_FONT, size, italic=False, color=INK_HEX)
        c.paragraph_format.space_after = Pt(10) if name == "FigCaption" else Pt(4)
        c.paragraph_format.space_before = Pt(4) if name == "FigCaption" else Pt(10)
        c.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == "FigCaption" else WD_ALIGN_PARAGRAPH.LEFT
        c.paragraph_format.line_spacing = 1.0
    code = _style(st, "CodeBlock", MONO_FONT, 8.5)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Cm(0.3)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _style(st, "Header", HEAD_FONT, 8.5, color=MUTED_HEX)
    _style(st, "Footer", HEAD_FONT, 9, color=MUTED_HEX)


def _hf_page_number(section, roman):
    section.footer.is_linked_to_previous = False
    section.header.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.text = ""
    run = hp.add_run("BunkWise: A Fuzzy-Logic (Mamdani) Attendance Decision-Support Tool  |  AISC Final Project Report")
    run.font.size = Pt(8.5)
    run.font.name = HEAD_FONT
    run.font.color.rgb = RGBColor.from_string(MUTED_HEX)
    _bottom_border(hp, color=LINE_HEX, sz=6, space=3)
    fp = section.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(fp, "PAGE", "1")
    sect_pr = section._sectPr
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), "lowerRoman" if roman else "decimal")
    pg.set(qn("w:start"), "1")
    _put(sect_pr, pg, SECTPR_ORDER)


def _clear_hf(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for part in (section.header, section.footer):
        for p in part.paragraphs:
            p.text = ""


def _heading(doc, level, text, page_break=False):
    p = doc.add_heading(text, level=level)
    if level == 1:
        _bottom_border(p, sz=6, space=3)
    if page_break:
        p.paragraph_format.page_break_before = True
    return p


def _table(doc, b, usable_cm):
    ncol = len(b["header"])
    cap = doc.add_paragraph(style="TabCaption")
    _add_runs(cap, f"**Table {b['label']}:** {b['caption']}")
    _keep_next(cap)
    t = doc.add_table(rows=1 + len(b["rows"]), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _cell_margins(t)
    widths = b["widths"]
    total = sum(widths)
    widths_cm = [w / total * usable_cm for w in widths]
    for ri, row in enumerate([b["header"]] + b["rows"]):
        tr = t.rows[ri]
        trpr = tr._tr.get_or_add_trPr()
        _put(trpr, OxmlElement("w:cantSplit"), TRPR_ORDER)
        if ri == 0:
            _put(trpr, OxmlElement("w:tblHeader"), TRPR_ORDER)
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            cell.width = Cm(widths_cm[ci])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            al = b["align"][ci]
            p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                           "r": WD_ALIGN_PARAGRAPH.RIGHT}[al]
            lines = str(val).split("\n")
            for li, line in enumerate(lines):
                if li:
                    p.add_run().add_break()
                if ri == 0:
                    _add_runs(p, f"**{line}**" if line else "", base_size=b["size"], font=HEAD_FONT, color="FFFFFF")
                elif ci in b["mono"]:
                    r = p.add_run(line)
                    r.font.name = MONO_FONT
                    _rfonts(r._r, MONO_FONT)
                    r.font.size = Pt(b["size"] - 0.5)
                else:
                    _add_runs(p, line, base_size=b["size"], font=HEAD_FONT)
            if ri == 0:
                _cell_shade(cell, ACCENT_HEX)
            elif ri % 2 == 0:
                _cell_shade(cell, "F5F4FC")
            edge = (4, LINE_HEX)
            _cell_borders(cell, top=edge, bottom=edge, left=edge, right=edge)
            cell.vertical_alignment = 1
    # keep the caption, header and first data row together; keep short tables whole
    keep_rows = t.rows[:-1] if len(b["rows"]) <= 10 else t.rows[:2]
    for tr in keep_rows:
        for c in tr.cells:
            for p in c.paragraphs:
                _keep_next(p)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _note(doc, b, usable_cm):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_margins(t, top=90, bottom=90, left=160, right=140)
    cell = t.rows[0].cells[0]
    cell.width = Cm(usable_cm)
    _cell_shade(cell, TINT_HEX)
    _cell_borders(cell, left=(24, ACCENT_HEX))
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(b["title"] + "  ")
    r.bold = True
    r.font.name = HEAD_FONT
    _rfonts(r._r, HEAD_FONT)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(ACCENT_HEX)
    _add_runs(p, b["text"], base_size=10.5)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    _put(trpr, OxmlElement("w:cantSplit"), TRPR_ORDER)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)


def write_docx(blocks, ctx, path, eq_dir):
    cover = ctx["cover"]
    doc = Document()
    _setup_styles(doc)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin, sec.right_margin = Cm(2.8), Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.1)
    usable_cm = 21.0 - 2.8 - 2.3
    usable_in = usable_cm / 2.54
    props = doc.core_properties
    props.title = cover["title"]
    props.author = "; ".join(n for n, _ in cover["students"])
    props.subject = cover["course"]
    props.keywords = "fuzzy logic, Mamdani, soft computing, Streamlit"
    props.comments = ""

    # ------------------------------------------------------------ cover page
    _clear_hf(sec)

    def cover_p(text, size, bold=False, color=INK_HEX, before=0, after=0, font=HEAD_FONT, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        r = p.add_run(text)
        r.font.size, r.bold, r.italic = Pt(size), bold, italic
        r.font.name = font
        _rfonts(r._r, font)
        r.font.color.rgb = RGBColor.from_string(color)
        return p

    cover_p("ARTIFICIAL INTELLIGENCE AND SOFT COMPUTING (AISC)", 11, True, MUTED_HEX, before=40, after=6)
    cover_p("Final Project Report", 11, False, MUTED_HEX, after=70)
    p = cover_p("BunkWise", 40, True, ACCENT_HEX, after=4)
    _bottom_border(p, sz=12, space=10)
    cover_p("A Fuzzy-Logic (Mamdani) Attendance Decision-Support Tool", 17, True, INK_HEX, before=14, after=10)
    cover_p("Skip the guesswork. Know your risk.", 11, False, MUTED_HEX, after=60, italic=True)
    cover_p("Course", 9.5, True, ACCENT_HEX, after=1)
    cover_p(cover["course"], 12, False, INK_HEX, after=26)
    cover_p("Submitted by", 9.5, True, ACCENT_HEX, after=6)
    for name, reg in cover["students"]:
        cover_p(name, 13, True, INK_HEX, after=0)
        cover_p(f"Registration Number {reg}", 11, False, MUTED_HEX, after=10)
    cover_p("", 6, after=14)
    for label, value in (("Submitted to", cover["faculty"]), ("Department / programme / institution", cover["institution"]),
                         ("Date of submission", cover["date"])):
        cover_p(label, 9, True, ACCENT_HEX, before=4, after=0)
        cover_p(value, 11, False, MUTED_HEX, after=4, italic=value.startswith("["))

    # ------------------------------------------------------------ front matter section
    s2 = doc.add_section(WD_SECTION.NEW_PAGE)
    _hf_page_number(s2, roman=True)

    def front_heading(text, first=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(0 if first else 18), Pt(10)
        if not first:
            p.paragraph_format.page_break_before = False
        r = p.add_run(text)
        r.bold, r.font.size, r.font.name = True, Pt(17), HEAD_FONT
        _rfonts(r._r, HEAD_FONT)
        r.font.color.rgb = RGBColor.from_string(ACCENT_HEX)
        _bottom_border(p, sz=6, space=3)
        _keep_next(p)
        return p

    front_heading("Table of Contents", first=True)
    p = doc.add_paragraph()
    _field(p, r'TOC \o "1-2" \h \z \u', "Right-click and choose Update Field to refresh the table of contents.")
    front_heading("List of Figures")
    p = doc.add_paragraph()
    _field(p, r'TOC \h \z \t "FigCaption,1"', "Right-click and choose Update Field to refresh this list.")
    front_heading("List of Tables")
    p = doc.add_paragraph()
    _field(p, r'TOC \h \z \t "TabCaption,1"', "Right-click and choose Update Field to refresh this list.")

    # ------------------------------------------------------------ body section
    s3 = doc.add_section(WD_SECTION.NEW_PAGE)
    _hf_page_number(s3, roman=False)
    s3.header.is_linked_to_previous = True
    s3.footer.is_linked_to_previous = True
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), "decimal")
    pg.set(qn("w:start"), "1")
    _put(s3._sectPr, pg, SECTPR_ORDER)

    eq_n = 0
    for idx, b in enumerate(blocks):
        t = b["t"]
        next_is_table = idx + 1 < len(blocks) and blocks[idx + 1]["t"] == "tab"
        if t == "h1":
            label = b["label"]
            _heading(doc, 1, f"{label} {b['title']}".strip(), page_break=b["pb"])
        elif t == "h2":
            _heading(doc, 2, f"{b['label']} {b['title']}".strip())
        elif t == "h3":
            _heading(doc, 3, b["title"])
        elif t == "p":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(p, b["text"])
            if next_is_table:
                _keep_next(p)  # a table's introduction stays on the page of its caption
        elif t in ("bul", "num"):
            style = "List Bullet" if t == "bul" else "List Number"
            for item in b["items"]:
                p = doc.add_paragraph(style=style)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(3)
                _add_runs(p, item)
            if t == "num":
                # restart numbering for every list: give each its own w:num instance
                _restart_numbering(doc, [pp for pp in doc.paragraphs[-len(b["items"]):]])
            doc.paragraphs[-1].paragraph_format.space_after = Pt(8)
        elif t in ("eq", "cases"):
            eq_n += 1
            img = os.path.join(eq_dir, f"eq_{eq_n:02d}.png")
            if t == "eq":
                render_equation(b["tex"], img)
            else:
                render_cases(b["lhs"], b["rows"], img)
            _equation_row(doc, img, b["label"], usable_cm)
        elif t == "fig":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(0)
            _keep_next(p)
            from PIL import Image
            with Image.open(b["path"]) as im:
                w_px, h_px = im.size
            width_in = min(b["width"], usable_in)
            height_in = width_in * h_px / w_px
            max_h = 8.3
            if height_in > max_h:
                width_in *= max_h / height_in
            p.add_run().add_picture(b["path"], width=Inches(width_in))
            cap = doc.add_paragraph(style="FigCaption")
            _add_runs(cap, f"**Figure {b['label']}:** {b['caption']}")
        elif t == "tab":
            _table(doc, b, usable_cm)
        elif t == "code":
            lines = b["text"].rstrip("\n").split("\n")
            p = doc.add_paragraph(style="CodeBlock")
            _para_shade(p, "F4F2EC")
            _left_border(p)
            p.paragraph_format.keep_together = True
            for i, line in enumerate(lines):
                if i:
                    p.add_run().add_break()
                r = p.add_run(line)
                r.font.name = MONO_FONT
                _rfonts(r._r, MONO_FONT)
        elif t == "note":
            _note(doc, b, usable_cm)
        elif t == "refs":
            for n, _, text in b["items"]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.0)
                p.paragraph_format.first_line_indent = Cm(-1.0)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _add_runs(p, f"[{n}]\t{text}", base_size=10)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
    doc.save(path)


def _left_border(p):
    ppr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    b = OxmlElement("w:left")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "18")
    b.set(qn("w:space"), "6")
    b.set(qn("w:color"), ACCENT_HEX)
    pb.append(b)
    _put(ppr, pb, PPR_ORDER)


def _equation_row(doc, img, label, usable_cm):
    from PIL import Image
    with Image.open(img) as im:
        w_px, h_px = im.size
    width_in = w_px / 300.0
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    side = 1.3
    mid = usable_cm - 2 * side
    for cell, w in zip(t.rows[0].cells, (side, mid, side)):
        cell.width = Cm(w)
        _cell_borders(cell)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        cell.vertical_alignment = 1
    c = t.rows[0].cells[1].paragraphs[0]
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    max_in = mid / 2.54
    if width_in > max_in:
        width_in = max_in
    c.add_run().add_picture(img, width=Inches(width_in))
    n = t.rows[0].cells[2].paragraphs[0]
    n.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = n.add_run(f"({label})")
    r.font.name = BODY_FONT
    r.font.size = Pt(11)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    _put(trpr, OxmlElement("w:cantSplit"), TRPR_ORDER)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.line_spacing = 0.5


def _restart_numbering(doc, paragraphs):
    """Make a numbered list restart at 1 by pointing its paragraphs at a fresh w:num."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    first = paragraphs[0]
    ppr = first._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    style_num_id = None
    if numpr is None:
        st = first.style
        sppr = st.element.find(qn("w:pPr"))
        if sppr is not None and sppr.find(qn("w:numPr")) is not None:
            style_num_id = sppr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
    else:
        style_num_id = numpr.find(qn("w:numId")).get(qn("w:val"))
    if style_num_id is None:
        return
    abstract_id = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == style_num_id:
            abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    if abstract_id is None:
        return
    ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    new_id = max(ids) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    an = OxmlElement("w:abstractNumId")
    an.set(qn("w:val"), abstract_id)
    num.append(an)
    ov = OxmlElement("w:lvlOverride")
    ov.set(qn("w:ilvl"), "0")
    so = OxmlElement("w:startOverride")
    so.set(qn("w:val"), "1")
    ov.append(so)
    num.append(ov)
    numbering.append(num)
    for p in paragraphs:
        ppr = p._p.get_or_add_pPr()
        np_ = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl")
        il.set(qn("w:val"), "0")
        ni = OxmlElement("w:numId")
        ni.set(qn("w:val"), str(new_id))
        np_.append(il)
        np_.append(ni)
        _put(ppr, np_, PPR_ORDER)
